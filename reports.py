import os
from docx import Document
from database import get_employees_sorted, get_recent_logs

async def generate_word_report():
    doc = Document()
    doc.add_heading("Smart Nazorat 2.0 - Oylik va Davomat Hisoboti", level=1)

    # 1-Bo'lim: Xodimlar ro'yxati
    doc.add_heading("1. Xodimlar va Stavkalar", level=2)
    employees = await get_employees_sorted()
    
    table1 = doc.add_table(rows=1, cols=4)
    hdr_cells = table1.rows[0].cells
    hdr_cells[0].text = 'ID'
    hdr_cells[1].text = 'Ism Familiya'
    hdr_cells[2].text = 'Bo\'lim'
    hdr_cells[3].text = 'Stavka (so\'m)'

    for emp in employees:
        row_cells = table1.add_row().cells
        row_cells[0].text = str(emp[0])
        row_cells[1].text = str(emp[1])
        row_cells[2].text = str(emp[2])
        row_cells[3].text = f"{emp[3]:,.0f}"

    # 2-Bo'lim: Keldi-ketdi vaqtlari
    doc.add_heading("2. Davomat Yozuvlari", level=2)
    logs = await get_recent_logs()

    table2 = doc.add_table(rows=1, cols=3)
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = 'Xodim'
    hdr_cells2[1].text = 'Harakat'
    hdr_cells2[2].text = 'Vaqt'

    for log in logs:
        row_cells = table2.add_row().cells
        row_cells[0].text = str(log[0])
        row_cells[1].text = "Keldi" if log[1] == "checked_in" else "Ketdi"
        row_cells[2].text = str(log[2])

    file_path = "Hisobot.docx"
    doc.save(file_path)
    return file_path
